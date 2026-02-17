from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any

from src.db import RecommendationDB


def _clamp_score(value: Any) -> float:
    try:
        score = float(value)
    except (TypeError, ValueError):
        score = 0.0
    return max(0.0, min(1.0, score))


def _parse_margin(value: Any) -> float:
    try:
        margin = float(value)
    except (TypeError, ValueError):
        return 0.0
    return margin / 100.0 if margin > 1.0 else margin


def _parse_bool(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def _migrate_csv(db: RecommendationDB, csv_path: Path) -> dict[str, int]:
    counts = {"total": 0, "active": 0, "pending_review": 0, "rejected": 0}

    with csv_path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            status = "active" if _parse_bool(row.get("llm_valid")) else "pending_review"
            payload = {
                "source_category": row.get("source_category", "").strip(),
                "target_sku": row.get("recommended_sku", "").strip(),
                "target_netsuite_id": (row.get("recommended_netsuite_id") or "").strip() or None,
                "target_name": row.get("recommended_name", "").strip(),
                "target_price": row.get("recommended_price") or None,
                "score": _clamp_score(row.get("llm_confidence")),
                "reasoning": row.get("llm_reason"),
                "relationship_type": row.get("relationship_type") or "complement",
                "status": status,
                "copurchase_count": row.get("enrichment_copurchase_count") or 0,
                "margin": _parse_margin(row.get("enrichment_margin")),
                "velocity": row.get("enrichment_velocity") or 0,
            }
            db.upsert_recommendation(payload)

            counts["total"] += 1
            counts[status] += 1

    return counts


def _migrate_xlsx(db: RecommendationDB, xlsx_path: Path) -> dict[str, int]:
    from openpyxl import load_workbook

    counts = {"total": 0, "active": 0, "pending_review": 0, "rejected": 0}
    wb = load_workbook(xlsx_path, read_only=True)
    try:
        if "Reviewed Pairings" not in wb.sheetnames:
            return counts

        ws = wb["Reviewed Pairings"]
        headers = [cell.value for cell in ws[1]]
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            source_category = (row.get("source_category") or "").strip()
            target_sku = (row.get("target_sku") or "").strip()
            if not source_category or not target_sku:
                continue

            keep = bool(row.get("keep", True))
            score = _clamp_score((row.get("relevance_score") or 0) / 5 if row.get("relevance_score") else 0.0)
            status = "active" if keep and score >= 0.4 else "pending_review"

            db.upsert_recommendation(
                {
                    "source_category": source_category,
                    "target_sku": target_sku,
                    "target_netsuite_id": None,
                    "target_name": row.get("target_category") or target_sku,
                    "target_price": None,
                    "score": score,
                    "reasoning": row.get("review_note"),
                    "relationship_type": row.get("relationship_type") or "complement",
                    "status": status,
                    "copurchase_count": row.get("copurchase_count") or 0,
                    "margin": _parse_margin(row.get("margin_pct")),
                    "velocity": row.get("sales_velocity") or 0,
                }
            )
            counts["total"] += 1
            counts[status] += 1
    finally:
        wb.close()

    return counts


def migrate_csv_to_db(csv_path: Path, xlsx_path: Path | None, db_path: Path) -> dict[str, int]:
    db = RecommendationDB(db_path)
    db.init_db()

    csv_counts = _migrate_csv(db, csv_path)
    xlsx_counts = {"total": 0, "active": 0, "pending_review": 0, "rejected": 0}
    if xlsx_path:
        xlsx_counts = _migrate_xlsx(db, xlsx_path)

    return {
        "total": csv_counts["total"] + xlsx_counts["total"],
        "active": csv_counts["active"] + xlsx_counts["active"],
        "pending_review": csv_counts["pending_review"] + xlsx_counts["pending_review"],
        "rejected": csv_counts["rejected"] + xlsx_counts["rejected"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Migrate recommendation CSV/XLSX data into SQLite DB")
    parser.add_argument("--csv", required=True, help="Path to full_recommendations_v2.csv")
    parser.add_argument("--xlsx", required=False, help="Optional path to dark_horse reviewed workbook")
    parser.add_argument("--db", required=True, help="Output SQLite database path")
    args = parser.parse_args()

    csv_path = Path(args.csv)
    xlsx_path = Path(args.xlsx) if args.xlsx else None
    db_path = Path(args.db)

    if not csv_path.exists():
        raise SystemExit(f"CSV not found: {csv_path}")
    if xlsx_path and not xlsx_path.exists():
        raise SystemExit(f"XLSX not found: {xlsx_path}")

    summary = migrate_csv_to_db(csv_path=csv_path, xlsx_path=xlsx_path, db_path=db_path)
    print(
        "Migrated {total} recommendations ({active} active, {pending_review} pending_review, {rejected} rejected)".format(
            **summary
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
