"""
CLI command: review — Post-process dark horse results with LLM scoring.

Reads raw dark horse Excel, sends pairings to a reviewer LLM in batches,
scores each pairing 1-5 on relevance, and outputs a cleaned workbook.

Usage:
    pre review data/dark_horse_discovery_glm_q8.xlsx --model athena
    pre review data/dark_horse_discovery_glm_q8.xlsx --model athena --min-score 3
    pre review data/dark_horse_discovery_glm_q8.xlsx --resume  # Resume interrupted run
"""

import sys
import typer
from pathlib import Path
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn, BarColumn, MofNCompleteColumn

from easyupsell.core.review import (
    load_pairings_from_excel,
    review_all_pairings,
    write_reviewed_excel,
)

app = typer.Typer(help="Review and score dark horse pairings using an LLM")
console = Console()

# Minimum margin threshold to qualify as "high priority"
HIGH_PRIORITY_MARGIN_THRESHOLD = 0.25


def _enrich_pairings(scored: list) -> list:
    """Attempt to enrich scored pairings with margin/velocity/copurchase data.
    
    Gracefully degrades to empty values if enrichment data is unavailable.
    Also sets the high_priority flag for items with high margin and zero copurchase.
    """
    enrichment = None
    try:
        sys.path.insert(0, str(Path(__file__).parent.parent.parent / "src"))
        from enrichment import EnrichmentService
        enrichment = EnrichmentService()
        if not enrichment.item_map and not enrichment.category_map:
            enrichment = None
    except Exception:
        pass  # Enrichment unavailable — degrade gracefully
    
    for p in scored:
        # Default enrichment values
        p.setdefault("target_sku", "")
        p.setdefault("sales_velocity", "")
        p.setdefault("margin_pct", "")
        p.setdefault("copurchase_count", "")
        
        # Try to look up enrichment data by category if available
        if enrichment:
            # Category-level lookup via category_map (iterate to find by name match)
            for cat_id, entries in enrichment.category_map.items():
                for entry in entries:
                    target_cat = p.get("target_category", "")
                    if entry.get("rec_category_name", "") == target_cat:
                        p["target_sku"] = entry.get("rec_sku", p["target_sku"])
                        p["sales_velocity"] = entry.get("rec_velocity", p["sales_velocity"])
                        p["margin_pct"] = entry.get("rec_margin", p["margin_pct"])
                        p["copurchase_count"] = entry.get("copurchase_count", p["copurchase_count"])
                        break
        
        # Set high_priority flag: high margin + zero copurchase = untapped opportunity
        try:
            margin = float(p.get("margin_pct", 0) or 0)
            copurchase = int(p.get("copurchase_count", 0) or 0)
            p["high_priority"] = margin >= HIGH_PRIORITY_MARGIN_THRESHOLD and copurchase == 0
        except (ValueError, TypeError):
            p["high_priority"] = False
    
    return scored


@app.callback(invoke_without_command=True)
def review(
    input_file: str = typer.Argument(..., help="Path to dark horse Excel file to review"),
    model: str = typer.Option("athena", help="LLM provider for review (athena, azure, etc.)"),
    batch_size: int = typer.Option(15, help="Pairings per review batch"),
    min_score: int = typer.Option(2, help="Minimum relevance score to keep (1-5)"),
    output: str = typer.Option(None, help="Output file path (default: adds _reviewed suffix)"),
    resume: bool = typer.Option(False, help="Resume from last checkpoint if interrupted"),
):
    """Review and score dark horse pairings using an LLM."""
    
    input_path = Path(input_file)
    if not input_path.exists():
        console.print(f"[red]File not found:[/red] {input_path}")
        raise typer.Exit(1)
    
    # Default output path
    if output:
        output_path = Path(output)
    else:
        output_path = input_path.with_stem(input_path.stem + "_reviewed")
    
    # Checkpoint path derived from output path
    checkpoint_path = output_path.with_stem(output_path.stem + "_checkpoint").with_suffix(".json")
    
    # 1. Load pairings
    console.print(f"\n[bold]Loading pairings from:[/bold] {input_path}")
    pairings = load_pairings_from_excel(input_path)
    console.print(f"  Found [cyan]{len(pairings)}[/cyan] pairings to review")
    
    if not pairings:
        console.print("[yellow]No pairings found. Nothing to review.[/yellow]")
        raise typer.Exit(0)
    
    # Check for existing checkpoint
    if resume and checkpoint_path.exists():
        console.print(f"  [yellow]⟳ Resuming from checkpoint:[/yellow] {checkpoint_path}")
    elif resume:
        console.print(f"  [dim]No checkpoint found — starting fresh[/dim]")
    
    # 2. Load gaps (pass through unchanged)
    from openpyxl import load_workbook
    gaps = []
    wb = load_workbook(input_path, read_only=True)
    if "Catalog Gaps" in wb.sheetnames:
        ws_gaps = wb["Catalog Gaps"]
        gap_headers = [cell.value for cell in ws_gaps[1]]
        for row in ws_gaps.iter_rows(min_row=2, values_only=True):
            record = dict(zip(gap_headers, row))
            if record.get("source_category"):
                gaps.append(record)
    wb.close()
    console.print(f"  Found [cyan]{len(gaps)}[/cyan] catalog gaps (pass-through)")
    
    # 3. Review pairings
    console.print(f"\n[bold]Reviewing with:[/bold] {model} (batch_size={batch_size}, min_score={min_score})")
    
    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(),
        MofNCompleteColumn(),
        console=console
    ) as progress:
        task = progress.add_task("Reviewing pairings...", total=len(pairings))
        
        def on_progress(done, total):
            progress.update(task, completed=done, description=f"Reviewed {done}/{total}...")
        
        scored = review_all_pairings(
            pairings,
            provider=model,
            batch_size=batch_size,
            progress_callback=on_progress,
            checkpoint_path=checkpoint_path,
            resume=resume,
        )
    
    # 4. Enrich with margin/velocity/copurchase data
    console.print(f"\n[bold]Enriching pairings...[/bold]")
    scored = _enrich_pairings(scored)
    
    priority_count = sum(1 for p in scored if p.get("high_priority"))
    if priority_count:
        console.print(f"  [gold1]★ High-priority opportunities:[/gold1] {priority_count}")
    
    # 5. Stats
    keeps = [p for p in scored if p.get("keep", True) and p.get("relevance_score", 0) >= min_score]
    rejects = [p for p in scored if not p.get("keep", True) or p.get("relevance_score", 0) < min_score]
    
    # Score distribution
    score_counts = {}
    for p in scored:
        s = p.get("relevance_score", -1)
        score_counts[s] = score_counts.get(s, 0) + 1
    
    console.print(f"\n[bold]Score Distribution:[/bold]")
    for score in sorted(score_counts.keys(), reverse=True):
        label = "★" * score if score > 0 else "?"
        bar = "█" * (score_counts[score] // 5 or 1)
        console.print(f"  {label:5s} ({score}): {score_counts[score]:4d} {bar}")
    
    console.print(f"\n  [green]✓ Keeping:[/green]  {len(keeps)}")
    console.print(f"  [red]✗ Rejected:[/red] {len(rejects)}")
    
    # 6. Write output
    stats = write_reviewed_excel(scored, gaps, output_path, min_score=min_score)
    
    console.print(f"\n[green]✓[/green] Saved reviewed results to [bold]{output_path}[/bold]")
    console.print(f"  Sheets: Reviewed Pairings ({stats['keeps']}), Rejected ({stats['rejects']}), Catalog Gaps ({stats['gaps']})")
