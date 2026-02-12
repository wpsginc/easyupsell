import csv
from pathlib import Path
from typing import List, Dict


def write_dark_horse_results(pairings: List[Dict], output_path: Path):
    """
    Write valid pairings to CSV.
    """
    if not pairings:
        return

    fieldnames = [
        "source_category",
        "target_category",
        "suggested_weight",
        "relationship_type",
        "reason",
        "llm_model_used",
        "concept_matched",
        "match_confidence"
    ]
    
    # Ensure directory exists
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    with open(output_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        
        for p in pairings:
            val = p.get("validation", {})
            writer.writerow({
                "source_category": p.get("source"),
                "target_category": p.get("target"),
                "suggested_weight": val.get("suggested_weight", 50),
                "relationship_type": val.get("relationship_type", "unknown"),
                "reason": val.get("reason", ""),
                "llm_model_used": "local",
                "concept_matched": p.get("concept"),
                "match_confidence": p.get("match_score")
            })

def write_catalog_gaps(gaps: List[Dict], output_path: Path):
    """
    Write catalog gaps to CSV.
    """
    if not gaps:
        return

    fieldnames = [
        "source_category",
        "missing_product_type",
        "suggested_weight",
        "reason"
    ]
    
    # Ensure directory exists
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    with open(output_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        
        for g in gaps:
            writer.writerow({
                "source_category": g.get("source"),
                "missing_product_type": g.get("missing_concept"),
                "suggested_weight": 0,
                "reason": "Concept not found in inventory"
            })


def write_dark_horse_excel(pairings: List[Dict], gaps: List[Dict], excel_path: Path, failures: List[Dict] = None):
    """
    Write Dark Horse results to a formatted Excel workbook with sheets:
    - 'Dark Horse Pairings' (hidden inventory we carry but isn't being bought together)
    - 'Catalog Gaps' (products we don't carry that would sell well)
    - 'Failures' (categories/concepts that failed LLM processing) — only if failures exist
    
    Same styling as the main recommendations workbook.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    if failures is None:
        failures = []

    excel_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # Style definitions (matching recommend_items_full.py)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, size=11, color="FFFFFF")
    pairing_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    gap_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fail_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")

    # ── Sheet 1: Dark Horse Pairings ──
    ws_pairings = wb.active
    ws_pairings.title = "Dark Horse Pairings"

    pairing_fields = [
        "source_category",
        "target_category",
        "concept_matched",
        "match_confidence",
        "relationship_type",
        "suggested_weight",
        "reason",
    ]

    # Headers
    for col_idx, field in enumerate(pairing_fields, 1):
        cell = ws_pairings.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data
    for row_idx, p in enumerate(pairings, 2):
        val = p.get("validation", {})
        row_data = {
            "source_category": p.get("source", ""),
            "target_category": p.get("target", ""),
            "concept_matched": p.get("concept", ""),
            "match_confidence": p.get("match_score", 0),
            "relationship_type": val.get("relationship_type", "unknown"),
            "suggested_weight": val.get("suggested_weight", 50),
            "reason": val.get("reason", ""),
        }
        for col_idx, field in enumerate(pairing_fields, 1):
            value = row_data.get(field, "")
            # Convert match_confidence to percentage
            if field == "match_confidence":
                try:
                    raw = float(value) if value else 0.0
                    value = raw / 100.0  # rapidfuzz scores are 0-100
                except (ValueError, TypeError):
                    value = 0.0
            elif field == "suggested_weight":
                try:
                    value = int(value) if value else 0
                except (ValueError, TypeError):
                    value = 0

            cell = ws_pairings.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = pairing_fill

            if field == "match_confidence":
                cell.number_format = '0%'

    # Auto-size columns
    for col_idx in range(1, len(pairing_fields) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws_pairings.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws_pairings.max_row + 1)
        )
        ws_pairings.column_dimensions[col_letter].width = min(max_len + 3, 50)

    # ── Sheet 2: Catalog Gaps ──
    ws_gaps = wb.create_sheet("Catalog Gaps")

    gap_fields = [
        "source_category",
        "missing_product_type",
    ]

    # Headers
    for col_idx, field in enumerate(gap_fields, 1):
        cell = ws_gaps.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data
    for row_idx, g in enumerate(gaps, 2):
        row_data = {
            "source_category": g.get("source", ""),
            "missing_product_type": g.get("missing_concept", ""),
        }
        for col_idx, field in enumerate(gap_fields, 1):
            cell = ws_gaps.cell(row=row_idx, column=col_idx, value=row_data.get(field, ""))
            cell.fill = gap_fill

    # Auto-size columns
    for col_idx in range(1, len(gap_fields) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws_gaps.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws_gaps.max_row + 1)
        )
        ws_gaps.column_dimensions[col_letter].width = min(max_len + 3, 50)

    # ── Sheet 3: Failures (only if any) ──
    if failures:
        ws_fail = wb.create_sheet("Failures")

        fail_fields = ["category", "stage", "error"]

        for col_idx, field in enumerate(fail_fields, 1):
            cell = ws_fail.cell(row=1, column=col_idx, value=field)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, f in enumerate(failures, 2):
            for col_idx, field in enumerate(fail_fields, 1):
                cell = ws_fail.cell(row=row_idx, column=col_idx, value=f.get(field, ""))
                cell.fill = fail_fill

        for col_idx in range(1, len(fail_fields) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(ws_fail.cell(row=r, column=col_idx).value or ""))
                for r in range(1, ws_fail.max_row + 1)
            )
            ws_fail.column_dimensions[col_letter].width = min(max_len + 3, 60)

    # Save
    wb.save(excel_path)
